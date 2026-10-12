
import gradio as gr
import pandas as pd
import re
from copy import copy
from openpyxl import load_workbook


# =========================================================
# 1. SPEED TESTING CLEANING
# =========================================================

REQUIRED_COLUMNS = [
    "Home Operator",
    "Application Throughput DL",
    "Application Throughput UL",
    "Latitude",
    "Longitude"
]


def get_file_path(file):
    return file if isinstance(file, str) else file.name


def clean_number(value):
    """Mengubah nilai throughput menjadi angka."""
    if pd.isna(value):
        return None

    value = str(value).strip().replace(" ", "")

    if not value:
        return None

    if "," in value and "." not in value:
        value = value.replace(",", ".")

    try:
        return float(value)
    except (ValueError, TypeError):
        return None


def process_file(file):
    """Membersihkan data Speed Testing."""
    if file is None:
        return None, "❌ Silakan upload file Excel terlebih dahulu."

    try:
        df = pd.read_excel(
            get_file_path(file),
            sheet_name="Speed_Testing",
            header=1
        )
    except ValueError:
        return None, "❌ Sheet 'Speed_Testing' tidak ditemukan."
    except Exception as e:
        return None, f"❌ Gagal membaca file: {e}"

    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]

    if missing:
        return None, (
            "❌ Kolom berikut tidak ditemukan:\n\n"
            + "\n".join(f"- {c}" for c in missing)
        )

    df_clean = df[REQUIRED_COLUMNS].copy()

    for col in [
        "Application Throughput DL",
        "Application Throughput UL"
    ]:
        df_clean[col] = df_clean[col].apply(clean_number)

    for col in ["Latitude", "Longitude"]:
        df_clean[col] = pd.to_numeric(
            df_clean[col], errors="coerce"
        )

    before = len(df_clean)

    # Hanya menghapus baris dengan koordinat tidak valid
    df_clean = df_clean.dropna(
        subset=["Latitude", "Longitude"]
    ).copy()

    removed = before - len(df_clean)
    output_file = "SpeedTesting_Cleaned.csv"

    try:
        df_clean.to_csv(
            output_file,
            index=False,
            encoding="utf-8-sig"
        )
    except Exception as e:
        return None, f"❌ Gagal menyimpan CSV: {e}"

    return output_file, (
        "✅ Speed Testing berhasil diproses!\n\n"
        f"Baris hasil: {len(df_clean)}\n"
        f"Jumlah kolom: {len(df_clean.columns)}\n"
        f"Baris tanpa koordinat valid yang dihapus: {removed}\n\n"
        f"File output: {output_file}"
    )


# =========================================================
# 2. COMPILE SUMMARY - SINGKATAN HEADER
# =========================================================

# Kode singkat untuk header kelompok
HEADER_MAP = {
    "FTP DL": "FTP DL",
    "FTP UL": "FTP UL",
    "Capacity DL": "CAP DL",
    "Capacity UL": "CAP UL",
    "SpeedTest": "ST",
    "Speed Test": "ST",
    "HTTP DL Transfer": "HTTP DL",
    "HTTP UL Transfer": "HTTP UL",
    "Browsing": "BROWSING",
    "Ping Test": "PING",
    "Youtube": "YT",
    "Instagram Post Picture": "IG POST PIC",
    "Instagram Post Video": "IG POST VID",
    "Whatsapp Call": "WA CALL",
    "Whatsapp Message": "WA MSG",
    "SMS Onnet": "SMS ONNET",
    "SMS Offnet": "SMS OFFNET",
    "Voice Call Onnet": "VCALL ONNET",
    "Voice Call Offnet": "VCALL OFFNET",
    "Average Voice Call": "AVG VOICE",
    "Signal Strenght 5G": "SIGS 5G",
    "Signal Strength 5G": "SIGS 5G",
    "Signal Qualty 5G": "SIGQ 5G",
    "Signal Quality 5G": "SIGQ 5G",
    "Signal Strenght 4G": "SIGS 4G",
    "Signal Strength 4G": "SIGS 4G",
    "Signal Qualty 4G": "SIGQ 4G",
    "Signal Quality 4G": "SIGQ 4G",
    "Signal Strenght 2G": "SIGS 2G",
    "Signal Strength 2G": "SIGS 2G",
    "Signal Qualty 2G": "SIGQ 2G",
    "Signal Quality 2G": "SIGQ 2G",
    "Bad Signal Strenght 5G": "BAD SIGS 5G",
    "Bad Signal Strength 5G": "BAD SIGS 5G",
    "Bad Signal Qualty 5G": "BAD SIGQ 5G",
    "Bad Signal Quality 5G": "BAD SIGQ 5G",
    "Bad Signal Strenght 4G": "BAD SIGS 4G",
    "Bad Signal Strength 4G": "BAD SIGS 4G",
    "Bad Signal Qualty 4G": "BAD SIGQ 4G",
    "Bad Signal Quality 4G": "BAD SIGQ 4G",
    "Bad Signal Strenght 2G": "BAD SIGS 2G",
    "Bad Signal Strength 2G": "BAD SIGS 2G",
    "Bad Signal Qualty 2G": "BAD SIGQ 2G",
    "Bad Signal Quality 2G": "BAD SIGQ 2G",
    "Detail Value 5G": "DETAIL Va5G",
    "Detail Value 4G": "DETAIL4G",
    "Detail Value 2G": "DETAIL2G",
}

# Urutkan agar nama kelompok yang panjang dicocokkan lebih dahulu
SORTED_GROUPS = sorted(
    HEADER_MAP.keys(),
    key=len,
    reverse=True
)


def normalize_text(value):
    """Menghilangkan spasi berlebih dan line break."""
    if value is None:
        return ""

    value = str(value).replace("\xa0", " ")
    value = value.replace("\n", " ")
    value = re.sub(r"\s+", " ", value).strip()

    if value.lower() in ("", "nan", "none"):
        return ""

    if value.startswith("Unnamed:"):
        return ""

    return value


def abbreviate_parent(parent):
    """Menyingkat nama kelompok berdasarkan kamus."""
    parent = normalize_text(parent)

    for original in SORTED_GROUPS:
        if parent.casefold() == original.casefold():
            return HEADER_MAP[original]

    # Jika belum ada di kamus, pertahankan nama yang sudah dirapikan
    return parent


def abbreviate_child(child, parent=""):
    """Meringkas subheader tanpa menghilangkan arti metrik."""
    child = normalize_text(child)
    parent = normalize_text(parent)

    if not child:
        return ""

    # Hapus awalan kelompok dari subheader jika berulang.
    # Contoh: "FTP UL Attempts" menjadi "Attempts".
    prefixes = [parent] if parent else []

    # Variasi penulisan kelompok di dalam subheader
    prefixes += [
        "IG Post Picture",
        "IG Post Video",
        "Instagram Post Picture",
        "Instagram Post Video",
        "FTP DL",
        "FTP UL",
        "Capacity DL",
        "Capacity UL",
        "Speed Test",
        "SpeedTest",
        "HTTP DL",
        "HTTP UL",
        "PING",
        "Youtube",
        "Whatsapp Call",
        "Whatsapp Message",
        "WA Call",
        "WA Message",
        "Browsing",
    ]

    for prefix in sorted(set(prefixes), key=len, reverse=True):
        if prefix and child.casefold().startswith(prefix.casefold()):
            child = child[len(prefix):].strip(" _-")

    # Normalisasi istilah dan satuan
    replacements = [
        (r"\bAttempts\b", "Attempts"),
        (r"\bAttemps\b", "Attempts"),
        (r"\bSuccess Attempts\b", "Success Attempts"),
        (r"\bSuccess Attemps\b", "Success Attempts"),
        (r"\bSuccess\b", "Success"),
        (r"\bAverage\b", "Avg"),
        (r"\bMedian\b", "Median"),
        (r"\bMaximum\b", "Max"),
        (r"\bMinimum\b", "Min"),
        (r"\bDuration\b", "Duration"),
        (r"\bSpeed\b", "Speed"),
        (r"\bLatency\b", "Latency"),
        (r"\bCategory\b", "Category"),
        (r"\bSample\b", "Sample"),
        (r"\bQuality\b", "Quality"),
        (r"\bStrenght\b", "Strength"),
        (r"\bStrength\b", "Strength"),
        # Satuan dipertahankan apa adanya dan selalu ditulis dalam tanda kurung.
        (r"\(\s*Mbps\s*\)|\bMbps\b", "(Mbps)"),
        (r"\(\s*ms\s*\)|\bms\b", "(ms)"),
        (r"\(\s*s\s*\)|\bs\b", "(s)"),
        (r"\(\s*%\s*\)|%", "(%)"),
        (r"\bMOS\b", "MOS"),
        (r"\bRTT\b", "RTT"),
        (r"\bRSRP\b", "RSRP"),
        (r"\bRSRQ\b", "RSRQ"),
        (r"\bSINR\b", "SINR"),
        (r"\bTTFP\b", "TTFP"),
        (r"\bCST\b", "CST"),
        (r"\bOnnet\b", "Onnet"),
        (r"\bOffnet\b", "Offnet"),
    ]

    for pattern, replacement in replacements:
        child = re.sub(pattern, replacement, child, flags=re.IGNORECASE)

    # Jadikan nama kolom ramah PostgreSQL
    child = re.sub(r"[^A-Za-z0-9 (%)-]+", " ", child)
    child = re.sub(r"\s+", " ", child).strip()

    return child


def make_unique_columns(columns):
    """Membuat setiap nama kolom unik."""
    result = []
    counts = {}

    for name in columns:
        base = name or "Column"
        counts[base] = counts.get(base, 0) + 1

        if counts[base] == 1:
            result.append(base)
        else:
            result.append(f"{base}_{counts[base]}")

    return result


def process_summary(file):
    """
    Menggabungkan dua baris header dan menyingkat nama kolom.
    Output tetap XLSX; data tidak difilter.
    """
    if file is None:
        return None, "❌ Silakan upload file Compile Summary terlebih dahulu."

    wb = None

    try:
        input_path = get_file_path(file)
        wb = load_workbook(input_path)

        # Gunakan sheet bernama Compile_Summary jika tersedia
        if "Compile_Summary" in wb.sheetnames:
            ws = wb["Compile_Summary"]
        else:
            ws = wb.active

        if ws.max_row < 2:
            return None, "❌ File harus memiliki dua baris header."

        max_col = ws.max_column

        parent_headers = [
            ws.cell(row=1, column=col).value
            for col in range(1, max_col + 1)
        ]

        child_headers = [
            ws.cell(row=2, column=col).value
            for col in range(1, max_col + 1)
        ]

        # Sebarkan header kelompok dari merged cells horizontal
        for merged_range in list(ws.merged_cells.ranges):
            if (
                merged_range.min_row == 1
                and merged_range.max_row == 1
                and merged_range.max_col > merged_range.min_col
            ):
                parent_value = ws.cell(
                    row=1,
                    column=merged_range.min_col
                ).value

                for col in range(
                    merged_range.min_col,
                    merged_range.max_col + 1
                ):
                    parent_headers[col - 1] = parent_value

        new_headers = []

        for index, (parent, child) in enumerate(
            zip(parent_headers, child_headers),
            start=1
        ):
            parent_text = normalize_text(parent)
            child_text = normalize_text(child)

            short_parent = abbreviate_parent(parent_text)
            short_child = abbreviate_child(child_text, parent_text)

            if short_parent and short_child:
                name = f"{short_parent}_{short_child}"
            elif short_child:
                name = short_child
            elif short_parent:
                name = short_parent
            else:
                name = f"Column_{index}"

            # Nama kolom tidak boleh diawali angka
            if name and name[0].isdigit():
                name = f"col_{name}"

            new_headers.append(name)

        new_headers = make_unique_columns(new_headers)

        # Simpan format header asli sebelum merge dilepas. Format subheader
        # biasanya memuat warna kelompok (FTP, CAP, dan seterusnya).
        header_styles = []
        for col in range(1, max_col + 1):
            parent_style = copy(ws.cell(row=1, column=col)._style)
            child_style = copy(ws.cell(row=2, column=col)._style)
            child_fill = ws.cell(row=2, column=col).fill

            if child_fill.fill_type:
                header_styles.append(child_style)
            else:
                header_styles.append(parent_style)

        # Simpan merge pada data agar posisinya dapat disesuaikan setelah
        # baris subheader dihapus. openpyxl tidak menggeser referensi merge
        # secara otomatis saat delete_rows() dipanggil.
        data_merges = []

        # Lepas merge yang beririsan dengan dua baris header.
        # Ini diperlukan agar header bisa ditulis menjadi satu baris.
        for merged_range in list(ws.merged_cells.ranges):
            if (
                merged_range.min_row <= 2
                and merged_range.max_row >= 1
            ):
                ws.unmerge_cells(str(merged_range))
            elif merged_range.min_row > 2:
                data_merges.append(
                    (
                        merged_range.min_row - 1,
                        merged_range.min_col,
                        merged_range.max_row - 1,
                        merged_range.max_col,
                    )
                )
                ws.unmerge_cells(str(merged_range))

        # Tulis header gabungan ke baris pertama
        for col, name in enumerate(new_headers, start=1):
            header_cell = ws.cell(row=1, column=col)
            header_cell.value = name
            header_cell._style = header_styles[col - 1]

        # Hapus baris header kedua; baris data di bawahnya tetap
        ws.delete_rows(2, 1)

        # Kembalikan merge data pada baris yang baru setelah header diringkas.
        for min_row, min_col, max_row, max_col in data_merges:
            ws.merge_cells(
                start_row=min_row,
                start_column=min_col,
                end_row=max_row,
                end_column=max_col,
            )

        output_file = "Compile_Summary_Cleaned.xlsx"
        wb.save(output_file)

        message = (
            "✅ Compile Summary berhasil diproses!\n\n"
            f"Jumlah kolom: {len(new_headers)}\n"
            "Format output: XLSX\n"
            "Header disatukan dan disingkat.\n"
            "Baris data tidak difilter.\n"
            "Nama kolom duplikat diberi akhiran angka.\n\n"
            f"File output: {output_file}"
        )

        return output_file, message

    except Exception as e:
        return None, f"❌ Gagal memproses Compile Summary: {e}"

    finally:
        if wb is not None:
            wb.close()


# =========================================================
# 3. TAMPILAN GRADIO
# =========================================================

css = """
body {
    background-color: #f5f7f9;
}

.gradio-container {
    max-width: 1000px !important;
    margin: auto;
}

.title {
    text-align: center;
    margin-bottom: 5px;
}

.subtitle {
    text-align: center;
    color: #6b7280;
    margin-bottom: 25px;
}

.footer {
    text-align: center;
    color: #9ca3af;
    font-size: 12px;
    margin-top: 25px;
}
"""


with gr.Blocks(title="Data Cleaning - Balmon Aceh") as demo:

    gr.HTML(
        """
        <div class="title">
            <h1>Data Cleaning</h1>
        </div>
        <div class="subtitle">
            Pengolahan data Speed Testing dan Compile Summary
        </div>
        """
    )

    # TAB SPEED TESTING
    with gr.Tab("Speed Testing"):

        gr.Markdown(
            """
            ### Cleaning Data Speed Testing

            Upload file Excel raw. Sistem membaca sheet
            `Speed_Testing` dengan header pada baris kedua.
            """
        )

        with gr.Row():
            with gr.Column():
                speed_input = gr.File(
                    label="Upload File Raw Speed Testing",
                    file_types=[".xlsm", ".xlsx", ".xls"],
                    type="filepath"
                )

                speed_button = gr.Button(
                    "⚙️ Proses Speed Testing",
                    variant="primary"
                )

            with gr.Column():
                speed_status = gr.Textbox(
                    label="Status Speed Testing",
                    placeholder="Hasil proses akan muncul di sini...",
                    lines=7,
                    interactive=False
                )

                speed_download = gr.File(
                    label="Download CSV Speed Testing",
                    file_types=[".csv"]
                )

        speed_button.click(
            fn=process_file,
            inputs=speed_input,
            outputs=[speed_download, speed_status]
        )

    # TAB COMPILE SUMMARY
    with gr.Tab("Compile Summary"):

        gr.Markdown(
            """
            ### Penyederhanaan Header Compile Summary

            Upload file Excel dengan dua baris header.
            Sistem menggabungkan dan menyingkat nama header
            menjadi nama kolom yang lebih ringkas untuk PostgreSQL.

            Output tetap berupa file Excel `.xlsx`.
            """
        )

        with gr.Row():
            with gr.Column():
                summary_input = gr.File(
                    label="Upload File Compile Summary",
                    file_types=[".xlsx", ".xlsm", ".xls"],
                    type="filepath"
                )

                summary_button = gr.Button(
                    "⚙️ Proses Compile Summary",
                    variant="primary"
                )

            with gr.Column():
                summary_status = gr.Textbox(
                    label="Status Compile Summary",
                    placeholder="Hasil proses akan muncul di sini...",
                    lines=7,
                    interactive=False
                )

                summary_download = gr.File(
                    label="Download Excel Compile Summary",
                    file_types=[".xlsx"]
                )

        summary_button.click(
            fn=process_summary,
            inputs=summary_input,
            outputs=[summary_download, summary_status]
        )

    gr.HTML(
        """
        <div class="footer">
            Data Cleaning • Balai Monitor Spektrum Frekuensi Radio Kelas II Aceh
        </div>
        """
    )


if __name__ == "__main__":
    demo.launch(css=css)
